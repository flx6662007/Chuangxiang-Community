"""从整理底稿生成审阅工作簿。可选依赖：pip install openpyxl；不参与数据库导入。"""
from __future__ import annotations

import argparse
import json
from collections import Counter
from datetime import date
from math import ceil
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.table import Table, TableStyleInfo
    from openpyxl.utils import get_column_letter
except ImportError as exc:
    raise SystemExit("仅生成工作簿需要 openpyxl：python -m pip install openpyxl") from exc

HERE = Path(__file__).resolve().parent
UNCERTAIN = {'2026052', '2026057', '2026063', '2026064', '2026067'}
EVIDENCE = {'page_read': '已读来源页面', 'official_index': '已确认索引或入口', 'search_result': '仅搜索结果，待核对'}
TYPES = {'rules': '规则与模板', 'problems': '赛题', 'course': '课程与指南', 'examples': '作品案例', 'tools': '工具', 'dataset': '数据集', 'code': '代码'}


def source_state(row):
    if row['code'] in UNCERTAIN:
        return '待确认对应范围'
    if any(s['verified_by'] in ('page_read', 'official_index') for s in row['competition_sources']):
        return '对应来源已有依据'
    return '来源对应尚待核验'


def parsed_date(value):
    if not value:
        return None
    try:
        return date.fromisoformat(value)
    except ValueError:
        return value


def add_sheet(book, name, headers, rows, widths, table_name=None, link_columns=()):
    sheet = book.create_sheet(name)
    sheet.sheet_view.showGridLines = False
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    for col, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(col)].width = width
    for row in sheet:
        for cell in row:
            cell.font = Font(name='Arial', size=10, color='25384B')
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            if isinstance(cell.value, date):
                cell.number_format = 'yyyy-mm-dd'
        if row[0].row == 1:
            for cell in row:
                cell.fill = PatternFill('solid', fgColor='DDECF9')
                cell.font = Font(name='Arial', size=10, bold=True, color='123F62')
                cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            sheet.row_dimensions[1].height = 32
        else:
            lines = max(sum(max(1, ceil(sum(2 if ord(ch) > 255 else 1 for ch in line) / max(6, widths[i] - 2))) for line in str(cell.value or '').split('\n')) for i, cell in enumerate(row))
            sheet.row_dimensions[row[0].row].height = min(409, max(36, lines * 14 + 8))
    if table_name:
        table = Table(displayName=table_name, ref=sheet.dimensions)
        table.tableStyleInfo = TableStyleInfo(name='TableStyleLight9', showRowStripes=True)
        sheet.add_table(table)
        sheet.freeze_panes = 'C2'
    for link_col, url_col in link_columns:
        for row in range(2, sheet.max_row + 1):
            url = sheet.cell(row, url_col).value
            if url:
                cell = sheet.cell(row, link_col, '打开来源')
                cell.hyperlink = url
                cell.font = Font(name='Arial', size=10, color='1666AE', underline='single')
    return sheet


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output', type=Path, default=HERE / '总索引.xlsx')
    args = parser.parse_args()
    payload = json.loads((HERE / '整理底稿.json').read_text(encoding='utf-8'))
    rows = sorted(payload['entries'], key=lambda item: item['code'])
    catalog_path = HERE.parents[2] / 'backend/competition_catalog/data/tongji-2026.json'
    catalog = {item['code']: item for item in json.loads(catalog_path.read_text(encoding='utf-8'))['entries']}
    assert [item['code'] for item in rows] == [str(2026000 + i) for i in range(1, 90)]
    assert all(item['name'] == catalog[item['code']]['name'] for item in rows)
    sources = [(row, source) for row in rows for source in row['competition_sources']]
    resources = [(row, resource) for row in rows for resource in row['learning_resources']]
    unread = [row['code'] for row in rows if not any(s['verified_by'] == 'page_read' for s in row['competition_sources'])]
    statuses = Counter(source_state(row) for row in rows)
    book = Workbook()
    book.remove(book.active)
    index = add_sheet(book, '赛事索引',
        ['序号','目录编号','赛事名称','目录等级','届次与范围','对应关系','赛事来源状态','来源数','学习资料数','主要入口','主要来源网址','赛事简介','缺项与后续核对','整理日期'],
        [[int(r['code'])-2026000,r['code'],r['name'],catalog[r['code']]['grade'],r['edition_note'],source_state(r),'赛事来源正文待补' if r['code'] in unread else '已读赛事来源页面',len(r['competition_sources']),len(r['learning_resources']),'',r['competition_sources'][0]['url'] if r['competition_sources'] else '',r['overview'],'\n'.join(r['gaps']),parsed_date(r['checked_on'])] for r in rows],
        [8,13,38,10,48,22,24,10,12,14,45,52,74,14], 'CompetitionIndex', [(10,11)])
    index.sheet_properties.tabColor = '2773AF'
    for row in range(2, index.max_row + 1):
        for col in (6, 7):
            if '待' in str(index.cell(row, col).value):
                index.cell(row, col).fill = PatternFill('solid', fgColor='FFF2CD')
    add_sheet(book, '学习资料',
        ['资料序号','目录编号','赛事名称','资料名称','资料类型','发布方','学习层次','适用范围','学习用途','获取条件','证据代码','核验程度','资料链接','来源网址','资料摘要','整理日期'],
        [[i,r['code'],r['name'],s['title'],TYPES.get(s['type'],s['type']),s['provider'],s['level'],s['applies_to'],s['why_useful'],s['access'],s['verified_by'],EVIDENCE[s['verified_by']],'',s['url'],s['summary'],parsed_date(r['checked_on'])] for i,(r,s) in enumerate(resources,1)],
        [10,13,36,44,16,30,12,42,60,20,20,24,14,50,60,14], 'LearningResources', [(13,14)])
    add_sheet(book, '赛事来源',
        ['来源序号','目录编号','赛事名称','页面名称','发布方','来源类型','证据代码','核验程度','原文发布日期','来源链接','来源网址','内容摘要','整理日期'],
        [[i,r['code'],r['name'],s['title'],s['provider'],{'official':'官方发布','university':'高校发布'}.get(s['source_type'],s['source_type']),s['verified_by'],EVIDENCE[s['verified_by']],parsed_date(s.get('published_on')),'',s['url'],s['summary'],parsed_date(r['checked_on'])] for i,(r,s) in enumerate(sources,1)],
        [10,13,36,44,32,14,20,24,16,14,52,72,14], 'CompetitionSources', [(10,11)])
    notes = [
        ['范围','目录第1—89项，编号2026001—2026089，重复目录编号按一项处理。'],
        ['资料性质','人工摘要与来源链接，不是网页、附件或视频全文归档。'],
        ['整理基准日',parsed_date(payload['checked_on'])],['赛事数',len(rows)],['赛事来源条目',len(sources)],['学习资料条目',len(resources)],
        ['对应来源已有依据',statuses['对应来源已有依据']],['待确认对应范围',statuses['待确认对应范围']],['来源对应尚待核验',statuses['来源对应尚待核验']],['赛事来源正文待补',len(unread)],
        ['page_read','已阅读链接页面，不代表其全部附件、视频和课程已阅读。'],['official_index','已确认索引或入口，资料内容待进一步核对。'],['search_result','仅搜索结果，直接访问受限或正文未读，不计为完成核验。'],
        ['状态口径','对应关系与页面是否已读是两个维度。official_index可提供赛事映射依据，但不表示已读正文。'],
        ['对应范围待确认编号','、'.join(sorted(UNCERTAIN))],['赛事来源正文待补编号','、'.join(unread)],
        ['届次规则','当届、下一届及往届按适用范围区分，不用旧日期推断当前报名状态。'],
        ['统计口径','资料条目数不等于独立文件数。同一个官方页面可能指向多份不同用途的资料。'],
        ['维护方法','先修改同目录整理底稿.json，再运行 python build_workbook.py。openpyxl仅为可选工作簿生成依赖，入库无需安装。'],
        ['使用方法','通过表头筛选编号、类型和证据代码，点击“打开来源”。空发布日期表示未找到，不推测填入。'],
        ['入库边界','不代表已经发布到前台、同步线上数据库或接通AI问答。']
    ]
    add_sheet(book, '说明', ['事项','说明'], notes, [29,110])
    args.output.parent.mkdir(parents=True, exist_ok=True)
    book.save(args.output)
    print(f'已写入 {args.output}：{len(rows)}项赛事、{len(sources)}条赛事来源、{len(resources)}条学习资料。')


if __name__ == '__main__':
    main()
